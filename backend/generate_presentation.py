from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image, ImageDraw, ImageFont

# 1. Create PPTX
prs = Presentation()
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)

title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "MedMate\nAI-Powered Triage & Medical Dossier Platform"
subtitle.text = "Submitted by: Rivanshu Gaur\nEnrollment Number: 03519051623\nSupervised by: Ms. Disha Dua\nMinor Project"

ppt_path = "/Users/rivanshu/Desktop/MedMate_Presentation.pptx"
prs.save(ppt_path)

# 2. Create JPEG Preview using Pillow
img = Image.new('RGB', (1920, 1080), color = (20, 30, 40)) # Dark teal/slate background
d = ImageDraw.Draw(img)

# Try to use a nice font, fallback to default
try:
    font_title = ImageFont.truetype("Arial", 90)
    font_sub = ImageFont.truetype("Arial", 40)
    font_small = ImageFont.truetype("Arial", 35)
except:
    font_title = ImageFont.load_default()
    font_sub = ImageFont.load_default()
    font_small = ImageFont.load_default()

# Title text
title_str = "MedMate"
sub_title_str = "AI-Powered Triage & Medical Dossier Platform"
d.text((960, 350), title_str, fill=(40, 200, 160), font=font_title, anchor="mm")
d.text((960, 470), sub_title_str, fill=(200, 220, 220), font=font_sub, anchor="mm")

# Details
details = "Minor Project\n\nSubmitted by: Rivanshu Gaur\nEnrollment Number: 03519051623\n\nSupervised by: Ms. Disha Dua"
d.text((960, 700), details, fill=(150, 180, 200), font=font_small, anchor="mm", align="center")

jpg_path = "/Users/rivanshu/Desktop/MedMate_Slide1_Preview.jpeg"
img.save(jpg_path)
print("SUCCESS")
